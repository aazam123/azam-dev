# Copyright (c) 2026, Hybrowlabs Pvt Ltd and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from openpyxl import Workbook
from frappe.utils.file_manager import save_file
from frappe.utils import now_datetime
import os
from openpyxl import load_workbook


class FREmployeeBenefitObligation(Document):
	def copy_child_table(self, source, target):
		self.get(target).clear()
		for row in source:
			self.append(target, row.as_dict())

	@frappe.whitelist()
	def fetch_templates(self):
		if not self.template:
			frappe.throw("Select the template first.")

		temp = frappe.get_doc("FR Employee Benefit Obligation Templates", self.template)
		self.copy_child_table(temp.amount_recognised_in_the_balance_sheet, "amount_recognised_in_the_balance_sheet")
		self.copy_child_table(temp.amount_recognised_in_statement_of_profit, "amount_recognised_in_statement_of_profit")
		self.copy_child_table(temp.reconciliation_of_defined_benefit_obligation, "reconciliation_of_defined_benefit_obligation")
		self.copy_child_table(temp.actuarial_assumptions, "actuarial_assumptions")
		self.copy_child_table(temp.experience_adjustments, "experience_adjustments")
		self.copy_child_table(temp.sensitivity_analysis, "sensitivity_analysis")
		self.copy_child_table(temp.maturity_analysis, "maturity_analysis")
		self.copy_child_table(temp.actuarial_assumption_for_compensated_absences, "actuarial_assumption_for_compensated_absences")
		self.copy_child_table(temp.amount_recognised_in_the_statement_of_profit_and_loss, "amount_recognised_in_the_statement_of_profit_and_loss")

@frappe.whitelist()
def download_excel(docname):
	import frappe
	from openpyxl import Workbook
	from openpyxl.styles import Font
	from io import BytesIO

	doc = frappe.get_doc("FR Employee Benefit Obligation", docname)

	table_map = [
		("amount_recognised_in_the_balance_sheet", "Amount Recognised in the Balance Sheet"),
		("amount_recognised_in_statement_of_profit", "Amount Recognised in Statement of Profit"),
		("reconciliation_of_defined_benefit_obligation", "Reconciliation of Defined Benefit Obligation"),
		("actuarial_assumptions", "Actuarial Assumptions"),
		("experience_adjustments", "Experience Adjustments"),
		("sensitivity_analysis", "Sensitivity Analysis"),
		("maturity_analysis", "Maturity Analysis"),
		("actuarial_assumption_for_compensated_absences", "Actuarial Assumption for Compensated Absences"),
		("amount_recognised_in_the_statement_of_profit_and_loss", "Amount Recognised in P&L"),
	]

	wb = Workbook()
	ws = wb.active
	ws.title = "Employee Benefit Obligation"

	row_cursor = 1
	bold = Font(bold=True)

	for fieldname, table_label in table_map:
		# Table title
		ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=2)
		cell = ws.cell(row=row_cursor, column=1, value=table_label)
		cell.font = bold
		row_cursor += 1

		# Table header
		ws.cell(row=row_cursor, column=1, value="Label").font = bold
		ws.cell(row=row_cursor, column=2, value="Value").font = bold
		row_cursor += 1

		# Table rows
		for row in doc.get(fieldname):
			ws.cell(row=row_cursor, column=1, value=row.label)
			ws.cell(row=row_cursor, column=2, value=row.value)
			row_cursor += 1

		# Space between tables
		row_cursor += 2

	xlsx_file = BytesIO()
	wb.save(xlsx_file)
	xlsx_file.seek(0)

	frappe.response.clear()
	frappe.response["filename"] = f"FR_Employee_Benefit_{docname}.xlsx"
	frappe.response["filecontent"] = xlsx_file.read()
	frappe.response["type"] = "binary"
	
@frappe.whitelist()
def upload_excel(docname, file_url):
	doc = frappe.get_doc("FR Employee Benefit Obligation", docname)
	file_path = frappe.get_site_path(file_url.lstrip("/"))

	table_map = {
		"Amount Recognised in the Balance Sheet": "amount_recognised_in_the_balance_sheet",
		"Amount Recognised in Statement of Profit": "amount_recognised_in_statement_of_profit",
		"Reconciliation of Defined Benefit Obligation": "reconciliation_of_defined_benefit_obligation",
		"Actuarial Assumptions": "actuarial_assumptions",
		"Experience Adjustments": "experience_adjustments",
		"Sensitivity Analysis": "sensitivity_analysis",
		"Maturity Analysis": "maturity_analysis",
		"Actuarial Assumption for Compensated Absences": "actuarial_assumption_for_compensated_absences",
		"Amount Recognised in P&L": "amount_recognised_in_the_statement_of_profit_and_loss",
	}

	wb = load_workbook(file_path)
	ws = wb.active

	current_table = None
	label_value_map = {}

	for row in ws.iter_rows(values_only=True):
		if row[0] in table_map:
			current_table = table_map[row[0]]
			label_value_map[current_table] = {}
			continue

		if row[0] == "Label":
			continue

		if not row[0]:
			current_table = None
			continue

		if current_table:
			label_value_map[current_table][row[0]] = row[1]

	# Update doc
	for fieldname, values in label_value_map.items():
		for row in doc.get(fieldname):
			if row.label in values:
				row.value = values[row.label]

	doc.save(ignore_permissions=True)
	frappe.db.commit()
